from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

import pymupdf
from docx import Document

from stages.analysis.compat import ContractError
from stages.analysis.inputs import ingest
from stages.analysis.profiling import MAX_CARD_BYTES, profile


class InputTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        (self.root / "input").mkdir()

    def test_text_markdown_stable_fragments(self) -> None:
        result = ingest("# Заголовок\n\nПервый абзац.\n\nВторой абзац.", [], self.root)
        self.assertEqual([source.id for source in result.sources], ["S-001", "S-002", "S-003"])
        self.assertEqual(result.sources[1].text, "Первый абзац.")

    def test_docx_paragraphs_and_table(self) -> None:
        doc = Document()
        doc.add_paragraph("Система должна хранить данные.")
        doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Заголовок"
        doc.save(self.root / "input/source.docx")
        result = ingest("", ["input/source.docx"], self.root)
        self.assertEqual(result.files[0].status, "parsed")
        self.assertEqual(len(result.sources), 2)

    def test_pdf_text_and_image_slide(self) -> None:
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), "Requirement one")
            pdf.new_page()
            pdf.save(self.root / "input/slides.pdf")
        result = ingest("", ["input/slides.pdf"], self.root)
        self.assertEqual(result.files[0].status, "partial")
        self.assertIn("2", result.files[0].reason)
        self.assertEqual(result.sources[0].text, "Requirement one")

    def test_image_missing_and_malformed_files(self) -> None:
        (self.root / "input/photo.png").write_bytes(b"image")
        (self.root / "input/bad.json").write_text("{broken")
        result = ingest("", ["input/photo.png", "input/missing.txt", "input/bad.json"], self.root)
        self.assertEqual([f.status for f in result.files], ["unparsed"] * 3)

    def test_outside_input_and_symlink_are_rejected(self) -> None:
        for path in ("../private.txt", "product/private.txt"):
            with self.subTest(path=path), self.assertRaises(ContractError):
                ingest("", [path], self.root)
        (self.root / "input/escape").symlink_to(self.root.parent, target_is_directory=True)
        with self.assertRaises(ContractError):
            ingest("", ["input/escape/private.txt"], self.root)

    def test_profile_json_fields_dates_anomalies_and_three_samples(self) -> None:
        rows = [{"id": i, "at": f"2026-10-{i+1:02d}T12:00:00+03:00", "value": None if i == 1 else i, "long": "x" * 4000} for i in range(3)]
        path = self.root / "input/data.json"
        path.write_text(json.dumps({"items": rows}))
        card = profile(path, "input/data.json")
        fields = card.data_schema["fields"]
        self.assertEqual(card.records, 3)
        self.assertEqual(card.data_schema["records_path"], "$.items")
        self.assertTrue(fields["id"]["all_unique"])
        self.assertEqual(fields["value"]["empty"], 1)
        self.assertIn("range", fields["at"])
        self.assertEqual(len(card.samples), 3)
        self.assertLess(len(card.model_dump_json().encode()), MAX_CARD_BYTES)

    def test_csv_delimiter_and_inconsistent_types(self) -> None:
        path = self.root / "input/data.csv"
        path.write_text("id;value\n1;0\n2;missing\n3;\n", encoding="utf-8")
        card = profile(path, path.name)
        self.assertEqual(card.records, 3)
        self.assertEqual(card.data_schema["fields"]["value"]["empty"], 1)
        self.assertTrue(card.anomalies)

    def test_geojson_line_is_not_invented_polygon(self) -> None:
        path = self.root / "input/data.geojson"
        path.write_text(json.dumps({"type": "FeatureCollection", "features": [{"type": "Feature", "properties": {"id": 1}, "geometry": {"type": "LineString", "coordinates": [[37.0, 55.0], [37.01, 55.0]]}}]}))
        card = profile(path, path.name)
        geometry = card.data_schema["geometry"]
        self.assertEqual(geometry["types"], {"LineString": 1})
        self.assertEqual(geometry["bbox"], [37.0, 55.0, 37.01, 55.0])
        self.assertGreater(geometry["lengths"]["sum"], 600)

    def test_wkt_text_ids_and_implicit_units(self) -> None:
        zones = [{"externalZoneId": str(517641 + i), "geometry": f"LINESTRING ({30 + i / 100} 59.9, {30.01 + i / 100} 59.9)",
                  "price": 28020, "priceDescription": "пн-пт 08:00-20:00, 280-280,2 ₽/час",
                  "tariff": {"periods": [{"from": 480, "to": 1200, "price": 56000}]}} for i in range(3)]
        zones[2]["geometry"] = "MULTILINESTRING ((30.1 59.9, 30.11 59.9), (30.2 59.9, 30.21 59.9))"
        path = self.root / "input/spb.json"
        path.write_text(json.dumps({"data": {"zones": zones}}, ensure_ascii=False), encoding="utf-8")
        schema = profile(path, path.name).data_schema
        self.assertEqual(schema["fields"]["externalZoneId"]["types"], {"numeric_string": 3})
        self.assertEqual(schema["fields"]["geometry"]["types"], {"wkt": 3})
        self.assertEqual(schema["geometry"]["types"], {"LineString": 2, "MultiLineString": 1})
        self.assertGreater(schema["geometry"]["lengths"]["min"], 500)
        self.assertEqual({path: hint["unit"] for path, hint in schema["units"].items()},
                         {"price": "kopecks", "tariff.periods[].price": "kopecks",
                          "tariff.periods[].from": "minutes_of_day", "tariff.periods[].to": "minutes_of_day"})
        csv_path = self.root / "input/data.csv"
        csv_path.write_text("id,value\n1,2\n", encoding="utf-8")
        self.assertEqual(profile(csv_path, csv_path.name).data_schema["fields"]["id"]["types"], {"int": 1})


if __name__ == "__main__":
    unittest.main()
