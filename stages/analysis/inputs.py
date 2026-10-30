"""S1: текстовые источники, PDF/DOCX и паспорта неразобранных файлов."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Literal

from pydantic import Field

from contracts.analysis import ContractModel, DataCard, Source
from stages.analysis.compat import ContractError
from stages.analysis.profiling import profile
from stages.analysis.storage import inside


class InputFile(ContractModel):
    file: str
    status: Literal["parsed", "partial", "unparsed"]
    reason: str = ""


class InputManifest(ContractModel):
    schema_version: Literal[1] = 1
    files: list[InputFile] = Field(default_factory=list)
    sources: list[Source] = Field(default_factory=list)
    noise: list[str] = Field(default_factory=list)
    data_cards: list[DataCard] = Field(default_factory=list)


def _paragraphs(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]


def ingest(requirements: str, attachments: list[str], workspace: Path) -> InputManifest:
    result = InputManifest()

    def add(texts: list[str], name: str) -> None:
        for text in texts:
            if text.strip():
                result.sources.append(Source(id=f"S-{len(result.sources) + 1:03d}", file=name, text=text))

    if requirements.strip():
        add(_paragraphs(requirements), "input/requirements.md")
        result.files.append(InputFile(file="input/requirements.md", status="parsed"))
    seen: set[Path] = set()
    for relative in attachments:
        path = inside(workspace, relative)
        if not path.is_relative_to(inside(workspace, "input")):
            raise ContractError("Приложение должно находиться внутри workspace/input")
        if path in seen or (path.name == "requirements.md" and requirements.strip()):
            continue
        seen.add(path)
        name = str(path.relative_to(workspace.resolve()))
        try:
            suffix = path.suffix.lower()
            if suffix in (".json", ".geojson", ".csv"):
                result.data_cards.append(profile(path, name))
            elif suffix in (".txt", ".md", ".markdown"):
                add(_paragraphs(path.read_text(encoding="utf-8-sig")), name)
            elif suffix == ".docx":
                from docx import Document
                document = Document(path)
                add([p.text for p in document.paragraphs if p.text.strip()], name)
                for table in document.tables:
                    add([" | ".join(cell.text for cell in row.cells) for row in table.rows], name)
                if document.inline_shapes:
                    result.files.append(InputFile(file=name, status="partial", reason="Изображения DOCX не разобраны"))
                    continue
            elif suffix == ".pdf":
                import pymupdf
                with pymupdf.open(path) as document:
                    missing: list[int] = []
                    text_pages = 0
                    for number, page in enumerate(document, 1):
                        text = page.get_text().strip()
                        if text:
                            add([text], name)
                            text_pages += 1
                        if not text or page.get_images():
                            missing.append(number)
                    if missing:
                        result.files.append(InputFile(file=name, status="partial" if text_pages else "unparsed", reason=f"Страницы без текста или с неразобранными изображениями: {missing}"))
                        continue
            else:
                result.files.append(InputFile(file=name, status="unparsed", reason="Формат или изображение пока не поддерживается"))
                continue
            result.files.append(InputFile(file=name, status="parsed"))
        except Exception as exc:
            result.files.append(InputFile(file=name, status="unparsed", reason=f"Не удалось разобрать файл ({type(exc).__name__})"))
    return result
